import traceback
from io import BytesIO
from flask import send_file
import re
from supabase import create_client
from uuid import uuid4
from secrets import token_urlsafe

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    Response,
    jsonify,
    abort
)

from datetime import date, timedelta, datetime

import os
import psycopg

from psycopg.rows import dict_row

from werkzeug.utils import secure_filename

from werkzeug.security import (
    check_password_hash,
    generate_password_hash
)

from functools import wraps


# ==================================================
# Flask
# ==================================================

app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "mono-admin-session-secret-2026"
)


# ==================================================
# 환경설정
# ==================================================

DATABASE_URL = os.environ.get(
    "DATABASE_URL"
)

SUPABASE_URL = os.environ.get(
    "SUPABASE_URL"
)

SUPABASE_SERVICE_ROLE_KEY = os.environ.get(
    "SUPABASE_SERVICE_ROLE_KEY"
)

SUPABASE_BUCKET = "mono-products"
BANK_ACCOUNT_NUMBER = os.environ.get("MONO_BANK_ACCOUNT_NUMBER", "").strip()

IS_VERCEL = bool(
    os.environ.get("VERCEL")
)


if IS_VERCEL:
    UPLOAD_FOLDER = "/tmp/mono-products"
else:
    UPLOAD_FOLDER = os.path.join(
        "static",
        "uploads",
        "products"
    )


os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER


# ==================================================
# 카테고리
# ==================================================

CATEGORIES = {
    "fashion": "패션의류/잡화",
    "food": "식품",
    "kitchen": "주방용품",
    "living": "생활용품",
    "hobby": "취미용품",
    "car": "자동차용품",
    "office": "사무용품",
    "sports": "스포츠용품"
}


# ==================================================
# DB 연결
# ==================================================

def get_db_connection():

    if not DATABASE_URL:
        raise RuntimeError(
            "DATABASE_URL 환경변수가 설정되어 있지 않습니다."
        )

    return psycopg.connect(
        DATABASE_URL,
        row_factory=dict_row,
        prepare_threshold=None
    )


# ==================================================
# Supabase Storage 연결
# ==================================================

def get_supabase_client():

    if not SUPABASE_URL:
        raise RuntimeError(
            "SUPABASE_URL 환경변수가 설정되어 있지 않습니다."
        )

    if not SUPABASE_SERVICE_ROLE_KEY:
        raise RuntimeError(
            "SUPABASE_SERVICE_ROLE_KEY 환경변수가 설정되어 있지 않습니다."
        )

    return create_client(
        SUPABASE_URL,
        SUPABASE_SERVICE_ROLE_KEY
    )


# ==================================================
# 상품 이미지 업로드
# ==================================================

def upload_product_image(
    uploaded_file,
    folder="product-images"
):

    if not uploaded_file:
        return None

    if not uploaded_file.filename:
        return None

    # 이미지 MIME 확인
    if not (
        uploaded_file.mimetype
        and uploaded_file.mimetype.startswith("image/")
    ):
        raise ValueError(
            "이미지 파일만 업로드할 수 있습니다."
        )

    # 안전한 파일명
    safe_name = secure_filename(
        uploaded_file.filename
    )

    # 확장자
    extension = os.path.splitext(
        safe_name
    )[1].lower()

    allowed_extensions = [
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
        ".gif"
    ]

    if extension not in allowed_extensions:
        raise ValueError(
            "JPG, JPEG, PNG, WEBP, GIF 이미지만 업로드할 수 있습니다."
        )

    # 랜덤 파일명
    filename = (
        str(uuid4())
        + extension
    )

    storage_path = (
        folder
        + "/"
        + filename
    )

    # 파일을 메모리 bytes로 읽기
    uploaded_file.stream.seek(0)

    file_data = uploaded_file.stream.read()

    if not file_data:
        raise ValueError(
            "업로드할 이미지 파일이 비어 있습니다."
        )

    supabase = get_supabase_client()

    # Supabase Storage 업로드
    supabase.storage.from_(
        SUPABASE_BUCKET
    ).upload(
        storage_path,
        file_data,
        {
            "content-type": uploaded_file.mimetype,
            "upsert": "false"
        }
    )

    # 공개 URL
    public_url = (
        supabase.storage
        .from_(SUPABASE_BUCKET)
        .get_public_url(storage_path)
    )

    return public_url


# ==================================================
# DB 생성 / 컬럼 업데이트
# ==================================================

def init_db():

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                CREATE TABLE IF NOT EXISTS products (
                    id SERIAL PRIMARY KEY,
                    name TEXT NOT NULL,
                    category TEXT NOT NULL,
                    image_url TEXT,
                    image_file TEXT,
                    description_html TEXT,
                    description_image BYTEA,
                    description_image_mime TEXT,
                    smartstore_price INTEGER,
                    smartstore_url TEXT,
                    coupang_price INTEGER,
                    coupang_url TEXT,
                    is_sold_out BOOLEAN DEFAULT FALSE,
                    created_at TIMESTAMP
                        DEFAULT CURRENT_TIMESTAMP
                )
            """)

            cur.execute("""
                ALTER TABLE products
                ADD COLUMN IF NOT EXISTS
                description_image BYTEA
            """)

            cur.execute("""
                ALTER TABLE products
                ADD COLUMN IF NOT EXISTS
                description_image_mime TEXT
            """)

            cur.execute("""
                ALTER TABLE products
                ADD COLUMN IF NOT EXISTS
                is_sold_out BOOLEAN
                DEFAULT FALSE
            """)

            cur.execute("""
                ALTER TABLE products
                ADD COLUMN IF NOT EXISTS bank_transfer_price INTEGER
            """)

            cur.execute("""
                CREATE TABLE IF NOT EXISTS product_options (
                    id SERIAL PRIMARY KEY,
                    product_id INTEGER NOT NULL,
                    option_name TEXT NOT NULL,
                    smartstore_price INTEGER,
                    smartstore_url TEXT,
                    coupang_price INTEGER,
                    coupang_url TEXT,

                    FOREIGN KEY (product_id)
                        REFERENCES products(id)
                        ON DELETE CASCADE
                )
            """)

            cur.execute("""
                ALTER TABLE product_options
                ADD COLUMN IF NOT EXISTS bank_transfer_price INTEGER
            """)

            cur.execute("""
                CREATE TABLE IF NOT EXISTS bank_orders (
                    id BIGSERIAL PRIMARY KEY,
                    product_id INTEGER,
                    product_name TEXT NOT NULL,
                    option_name TEXT NOT NULL,
                    depositor_name VARCHAR(50) NOT NULL,
                    amount INTEGER NOT NULL CHECK (amount > 0),
                    status TEXT NOT NULL DEFAULT 'reported'
                        CHECK (status IN ('reported','paid','cancelled','refunded')),
                    created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    paid_at TIMESTAMPTZ,
                    refunded_at TIMESTAMPTZ
                )
            """)
            # 신규 주문 수량을 별도 저장 (기존 주문은 옵션명 끝의 × 수량으로 집계)
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS quantity INTEGER")
            # 주문 시점의 배송지 스냅샷 (기존 주문은 NULL 유지)
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS recipient_name VARCHAR(100)")
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS recipient_phone VARCHAR(30)")
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS postal_code VARCHAR(10)")
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS address TEXT")
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS address_detail TEXT")
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS user_id INTEGER")

            # 포인트 결제: 주문별 원금/사용 포인트/실제 입금액을 보관
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS point_used INTEGER NOT NULL DEFAULT 0")
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS gross_amount INTEGER")
            cur.execute("ALTER TABLE bank_orders ADD COLUMN IF NOT EXISTS order_group_id UUID")
            cur.execute("UPDATE bank_orders SET gross_amount = amount WHERE gross_amount IS NULL")
            # 포인트 100% 결제는 입금액이 0원일 수 있다.
            cur.execute("ALTER TABLE bank_orders DROP CONSTRAINT IF EXISTS bank_orders_amount_check")
            cur.execute("ALTER TABLE bank_orders ADD CONSTRAINT bank_orders_amount_check CHECK (amount >= 0)")
            cur.execute("ALTER TABLE bank_orders ADD CONSTRAINT bank_orders_point_nonnegative CHECK (point_used >= 0)" ) if False else None
            cur.execute("CREATE INDEX IF NOT EXISTS idx_bank_orders_group ON bank_orders(order_group_id)")

            # 회원별 장바구니: 브라우저 세션이 아닌 DB에 회원 ID로 저장
            cur.execute("""CREATE TABLE IF NOT EXISTS member_cart_items (
                id BIGSERIAL PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                product_id INTEGER NOT NULL REFERENCES products(id) ON DELETE CASCADE,
                option_id INTEGER REFERENCES product_options(id) ON DELETE CASCADE,
                quantity INTEGER NOT NULL CHECK (quantity BETWEEN 1 AND 20),
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )""")
            cur.execute("""CREATE UNIQUE INDEX IF NOT EXISTS idx_member_cart_unique
                ON member_cart_items (user_id, product_id, COALESCE(option_id, 0))""")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_member_cart_user ON member_cart_items(user_id)")

            # 기존 회원 정보를 유지하면서 배송지 컬럼 추가
            cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS postal_code VARCHAR(10)")
            cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS address TEXT")
            cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS address_detail TEXT")

            cur.execute("CREATE INDEX IF NOT EXISTS idx_bank_orders_status ON bank_orders(status)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_bank_orders_paid_at ON bank_orders(paid_at)")

        conn.commit()

    finally:

        conn.close()


# ==================================================
# 서버 시작 후 DB 초기화
# ==================================================

_db_initialized = False


@app.before_request
def require_member_for_purchase():
    """가격/주문 데이터는 비회원에게 제공하지 않는다."""
    protected = {"view_cart", "add_to_cart", "update_cart", "checkout", "create_bank_order"}
    if request.endpoint in protected and not session.get("user_logged_in"):
        if request.path.startswith("/api/"):
            return jsonify(error="회원 로그인 후 주문할 수 있습니다."), 401
        return redirect(url_for("login"))


@app.before_request
def ensure_database():

    global _db_initialized

    if not _db_initialized:
        init_db()
        _db_initialized = True




def mask_public_prices(products):
    """상품 이름/이미지는 공개하되 가격 필드는 템플릿에 전달하지 않는다."""
    masked = []
    for item in products:
        p = dict(item)
        for key in list(p):
            if "price" in key.lower():
                p[key] = None
        p["listing_options"] = [
            {k: (None if "price" in k.lower() else v) for k, v in dict(opt).items()}
            for opt in p.get("listing_options", [])
        ]
        masked.append(p)
    return masked


def attach_listing_options(cur, products):
    """상품 목록에 옵션별 계좌이체 가격을 한 번의 조회로 연결."""
    if not products:
        return
    product_ids = [p["id"] for p in products]
    cur.execute("""
        SELECT id, product_id, option_name, bank_transfer_price
        FROM product_options
        WHERE product_id = ANY(%s)
        ORDER BY id ASC
    """, (product_ids,))
    by_product = {pid: [] for pid in product_ids}
    for option in cur.fetchall():
        by_product[option["product_id"]].append(option)
    for product in products:
        product["listing_options"] = by_product[product["id"]]

# ==================================================
# 메인
# ==================================================

@app.route("/")
def home():
    # 결제 완료된 계좌이체 주문의 상품별 실제 판매 수량 순위.
    # 구형 주문은 option_name 끝의 ' × N' 패턴을 이용한다.
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""
                WITH sold AS (
                    SELECT product_id,
                           SUM(COALESCE(quantity,
                               CASE WHEN option_name ~ ' × [0-9]+$'
                                    THEN (substring(option_name FROM ' × ([0-9]+)$'))::INTEGER
                                    ELSE 1 END)) AS units,
                           MAX(paid_at) AS last_paid
                    FROM bank_orders
                    WHERE status = 'paid' AND product_id IS NOT NULL
                    GROUP BY product_id
                )
                SELECT p.* FROM sold s
                JOIN products p ON p.id = s.product_id
                ORDER BY s.units DESC, s.last_paid DESC NULLS LAST, p.id DESC
                LIMIT 20
            """)
            products = cur.fetchall()
            attach_listing_options(cur, products)
    finally:
        conn.close()
    if not session.get("user_logged_in"):
        products = mask_public_prices(products)
    return render_template("index.html", products=products, home_top20=True)


@app.route("/events/<event_key>")
def event_page(event_key):
    pages = {
        "purchase-review": ("구매 인증", "구매 후기·인증 기능을 준비 중입니다."),
        "roulette": ("랜덤 룰렛", "룰렛 이벤트를 준비 중입니다. 아직 참여하거나 포인트를 받을 수 없습니다."),
        "support": ("고객센터", "문의 및 상담 접수 기능을 준비 중입니다."),
        "specials": ("특가 상품", "특가 상품 등록 기능을 준비 중입니다."),
    }
    if event_key not in pages:
        abort(404)
    title, message = pages[event_key]
    return render_template("event_page.html", title=title, message=message)


# ==================================================
# 전체 상품
# ==================================================

@app.route("/products")
def all_products():

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT *
                FROM products
                ORDER BY id DESC
            """)

            products = cur.fetchall()
            attach_listing_options(cur, products)

    finally:
        conn.close()

    if not session.get("user_logged_in"):
        products = mask_public_prices(products)

    return render_template(
        "index.html",
        products=products, home_top20=False
    )


# ==================================================
# 카테고리
# ==================================================

@app.route("/category/<category_slug>")
def category_page(category_slug):

    category_name = CATEGORIES.get(
        category_slug
    )

    if not category_name:
        return redirect(
            url_for("home")
        )

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT *
                FROM products
                WHERE category = %s
                ORDER BY id DESC
            """, (
                category_name,
            ))

            products = cur.fetchall()
            attach_listing_options(cur, products)

    finally:
        conn.close()

    if not session.get("user_logged_in"):
        products = mask_public_prices(products)

    return render_template(
        "category.html",
        products=products,
        category_name=category_name,
        category_slug=category_slug
    )


# ==================================================
# 상품 상세
# ==================================================

def get_order_csrf():
    if "order_csrf" not in session:
        session["order_csrf"] = token_urlsafe(32)
    return session["order_csrf"]


@app.route("/product/<int:product_id>")
def product_detail(product_id):

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    name,
                    category,
                    image_url,
                    image_file,
                    description_html,

                    CASE
                        WHEN description_image IS NOT NULL
                        THEN TRUE
                        ELSE FALSE
                    END AS has_description_image,

                    smartstore_price,
                    smartstore_url,
                    coupang_price,
                    coupang_url,
                    bank_transfer_price,

                    is_sold_out,

                    created_at

                FROM products
                WHERE id = %s
            """, (
                product_id,
            ))

            product = cur.fetchone()

            if product is None:

                return (
                    "상품을 찾을 수 없습니다.",
                    404
                )

            cur.execute("""
                SELECT *
                FROM product_options
                WHERE product_id = %s
                ORDER BY id ASC
            """, (
                product_id,
            ))

            options = cur.fetchall()

    finally:

        conn.close()

    if not session.get("user_logged_in"):
        product = {k: (None if "price" in k.lower() else v) for k, v in dict(product).items()}
        options = [
            {k: (None if "price" in k.lower() else v) for k, v in dict(opt).items()}
            for opt in options
        ]

    return render_template(
        "product_detail.html",
        product=product,
        options=options,
        bank_account_number=BANK_ACCOUNT_NUMBER,
        order_csrf=get_order_csrf()
    )


# ==================================================
# 상품 설명 이미지 출력
# ==================================================

@app.route(
    "/product/<int:product_id>/description-image"
)
def product_description_image(product_id):

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    description_image,
                    description_image_mime
                FROM products
                WHERE id = %s
            """, (
                product_id,
            ))

            product = cur.fetchone()

    finally:
        conn.close()

    if (
        product is None
        or product["description_image"] is None
    ):
        return (
            "상품 설명 이미지를 찾을 수 없습니다.",
            404
        )

    image_data = bytes(
        product["description_image"]
    )

    image_mime = (
        product["description_image_mime"]
        or "image/jpeg"
    )

    return Response(
        image_data,
        mimetype=image_mime
    )


# ==================================================
# 회원가입 아이디 중복 확인
# ==================================================

@app.route("/check_username", methods=["POST"])
def check_username():
    data = request.get_json(silent=True) or {}
    if not session.get("register_csrf") or data.get("csrf_token") != session.get("register_csrf"):
        return jsonify(ok=False, message="페이지를 새로고침해 주세요."), 403

    username = str(data.get("username") or "").strip()
    session.pop("checked_register_username", None)
    if not username:
        return jsonify(ok=False, message="아이디를 입력해 주세요."), 400
    if len(username) > 50:
        return jsonify(ok=False, message="아이디는 50자 이내로 입력해 주세요."), 400

    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM users WHERE username = %s LIMIT 1", (username,))
            exists = cur.fetchone() is not None
    finally:
        conn.close()

    if exists:
        return jsonify(ok=True, available=False, message="이미 사용 중인 아이디입니다.")
    session["checked_register_username"] = username
    return jsonify(ok=True, available=True, message="사용 가능한 아이디입니다.")


# ==================================================
# 일반 회원가입
# ==================================================

@app.route(
    "/register",
    methods=["GET", "POST"]
)
def register():

    if "register_csrf" not in session:
        session["register_csrf"] = token_urlsafe(32)

    error = None

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        username = request.form.get(
            "username",
            ""
        ).strip()

        phone = request.form.get(
            "phone",
            ""
        ).strip()

        postal_code = request.form.get("postal_code", "").strip()
        address = request.form.get("address", "").strip()
        address_detail = request.form.get("address_detail", "").strip()

        email = request.form.get(
            "email",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        password_confirm = request.form.get(
            "password_confirm",
            ""
        )

        if not name:
            error = "이름을 입력해주세요."

        elif not username:
            error = "아이디를 입력해주세요."
        elif len(username) > 50:
            error = "아이디는 50자 이내로 입력해주세요."
        elif request.form.get("register_csrf") != session.get("register_csrf"):
            error = "페이지를 새로고침한 후 다시 시도해 주세요."
        elif session.get("checked_register_username") != username:
            error = "아이디 중복 확인을 먼저 진행해 주세요."

        elif not postal_code or not address:
            error = "우편번호와 배송지 주소를 입력해주세요."
        elif not re.fullmatch(r"[0-9]{5}", postal_code):
            error = "우편번호는 5자리 숫자로 입력해주세요."
        elif len(address) > 300 or len(address_detail) > 300:
            error = "주소는 각 300자 이내로 입력해주세요."
        elif not email:
            error = "이메일을 입력해주세요."

        elif not password:
            error = "비밀번호를 입력해주세요."

        elif password != password_confirm:
            error = "비밀번호가 일치하지 않습니다."

        if error is None:

            conn = get_db_connection()

            try:

                with conn.cursor() as cur:

                    cur.execute("""
                        SELECT id
                        FROM users
                        WHERE username = %s
                    """, (
                        username,
                    ))

                    existing_username = cur.fetchone()

                    if existing_username:

                        error = (
                            "이미 사용 중인 아이디입니다."
                        )

                    else:

                        cur.execute("""
                            SELECT id
                            FROM users
                            WHERE email = %s
                        """, (
                            email,
                        ))

                        existing_email = cur.fetchone()

                        if existing_email:

                            error = (
                                "이미 가입된 이메일입니다."
                            )

                        else:

                            password_hash = (
                                generate_password_hash(
                                    password
                                )
                            )

                            cur.execute("""
                                INSERT INTO users (
                                    name,
                                    username,
                                    phone,
                                    email,
                                    password_hash,
                                    postal_code,
                                    address,
                                    address_detail
                                )
                                VALUES (
                                    %s,
                                    %s,
                                    %s,
                                    %s,
                                    %s,
                                    %s,
                                    %s,
                                    %s
                                )
                            """, (
                                name,
                                username,
                                phone,
                                email,
                                password_hash,
                                postal_code,
                                address,
                                address_detail
                            ))

                            conn.commit()

            except Exception:

                conn.rollback()
                raise

            finally:
                conn.close()

            if error is None:
                session.pop("checked_register_username", None)
                session.pop("register_csrf", None)
                return redirect(
                    url_for("login")
                )

    return render_template(
        "register.html",
        error=error,
        register_csrf=session["register_csrf"]
    )


# ==================================================
# 일반 사용자 로그인
# ==================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    error = None

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        conn = get_db_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT *
                    FROM users
                    WHERE username = %s
                """, (
                    username,
                ))

                user = cur.fetchone()

        finally:
            conn.close()

        if (
            user
            and check_password_hash(
                user["password_hash"],
                password
            )
        ):

            session.pop("mono_cart", None)
            session["user_logged_in"] = True
            session["user_id"] = user["id"]
            session["username"] = user["username"]

            return redirect(
                url_for("home")
            )

        error = (
            "아이디 또는 비밀번호가 올바르지 않습니다."
        )

    return render_template(
        "login.html",
        error=error
    )


# ==================================================
# 일반 사용자 로그아웃
# ==================================================

@app.route("/logout")
def logout():

    session.pop(
        "user_logged_in",
        None
    )

    session.pop(
        "user_id",
        None
    )

    session.pop(
        "username",
        None
    )

    return redirect(
        url_for("home")
    )


# ==================================================
# 내 정보
# ==================================================

@app.route("/profile")
def profile():

    if not session.get(
        "user_logged_in"
    ):
        return redirect(
            url_for("login")
        )

    user_id = session.get(
        "user_id"
    )

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    name,
                    username,
                    phone,
                    email,
                    postal_code,
                    address,
                    address_detail,
                    level,
                    point,
                    exp,
                    streak,
                    last_attendance,
                    created_at
                FROM users
                WHERE id = %s
            """, (
                user_id,
            ))

            user = cur.fetchone()

    finally:
        conn.close()

    if user is None:

        session.pop(
            "user_logged_in",
            None
        )

        session.pop(
            "user_id",
            None
        )

        session.pop(
            "username",
            None
        )

        return redirect(
            url_for("login")
        )

    # 주문은 반드시 현재 로그인한 회원 ID로 제한한다.
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""
                SELECT id, order_group_id, product_name, option_name,
                       gross_amount, point_used, amount, status, created_at
                FROM bank_orders
                WHERE user_id = %s
                ORDER BY created_at DESC, id DESC
                LIMIT 100
            """, (user_id,))
            order_rows = cur.fetchall()
    finally:
        conn.close()

    # 한 번의 장바구니 결제로 생성된 여러 상품 행을 주문 1건으로 표시한다.
    orders_by_key = {}
    for row in order_rows:
        group_key = str(row["order_group_id"]) if row["order_group_id"] else f'legacy-{row["id"]}'
        if group_key not in orders_by_key:
            orders_by_key[group_key] = {
                "order_id": row["id"],
                "group_id": group_key,
                "created_at": row["created_at"],
                "gross_amount": 0,
                "point_used": 0,
                "amount": 0,
                "items": [],
                "statuses": set(),
            }
        order = orders_by_key[group_key]
        order["gross_amount"] += row["gross_amount"] if row["gross_amount"] is not None else row["amount"]
        order["point_used"] += row["point_used"] or 0
        order["amount"] += row["amount"]
        order["items"].append({"id": row["id"], "name": row["product_name"],
                               "option": row["option_name"], "status": row["status"]})
        order["statuses"].add(row["status"])
    orders = list(orders_by_key.values())
    for order in orders:
        order["status"] = next(iter(order["statuses"])) if len(order["statuses"]) == 1 else "mixed"
        del order["statuses"]

    return render_template(
        "profile.html",
        user=user,
        orders=orders
    )


# ==================================================
# 내 정보 수정
# ==================================================

@app.route(
    "/profile/edit",
    methods=["GET", "POST"]
)
def edit_profile():

    if not session.get(
        "user_logged_in"
    ):
        return redirect(
            url_for("login")
        )

    user_id = session.get(
        "user_id"
    )

    error = None

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    name,
                    username,
                    phone,
                    email,
                    postal_code,
                    address,
                    address_detail,
                    password_hash
                FROM users
                WHERE id = %s
            """, (
                user_id,
            ))

            user = cur.fetchone()

        if user is None:

            session.pop(
                "user_logged_in",
                None
            )

            session.pop(
                "user_id",
                None
            )

            session.pop(
                "username",
                None
            )

            return redirect(
                url_for("login")
            )

        if request.method == "POST":

            name = request.form.get(
                "name",
                ""
            ).strip()

            phone = request.form.get(
                "phone",
                ""
            ).strip()

            email = request.form.get(
                "email",
                ""
            ).strip()

            postal_code = request.form.get("postal_code", "").strip()
            address = request.form.get("address", "").strip()
            address_detail = request.form.get("address_detail", "").strip()

            current_password = request.form.get(
                "current_password",
                ""
            )

            if not name:
                error = "이름을 입력해주세요."

            elif not postal_code or not address:
                error = "우편번호와 배송지 주소를 입력해주세요."
            elif not re.fullmatch(r"[0-9]{5}", postal_code):
                error = "우편번호는 5자리 숫자로 입력해주세요."
            elif len(address) > 300 or len(address_detail) > 300:
                error = "주소는 각 300자 이내로 입력해주세요."
            elif not email:
                error = "이메일을 입력해주세요."

            elif not current_password:
                error = "현재 비밀번호를 입력해주세요."

            elif not check_password_hash(
                user["password_hash"],
                current_password
            ):
                error = (
                    "비밀번호가 올바르지 않습니다."
                )

            if error is None:

                with conn.cursor() as cur:

                    cur.execute("""
                        SELECT id
                        FROM users
                        WHERE email = %s
                        AND id != %s
                    """, (
                        email,
                        user_id
                    ))

                    existing_email = cur.fetchone()

                if existing_email:
                    error = (
                        "이미 사용 중인 이메일입니다."
                    )

            if error is None:

                with conn.cursor() as cur:

                    cur.execute("""
                        UPDATE users
                        SET
                            name = %s,
                            phone = %s,
                            email = %s,
                            postal_code = %s,
                            address = %s,
                            address_detail = %s
                        WHERE id = %s
                    """, (
                        name,
                        phone,
                        email,
                        postal_code,
                        address,
                        address_detail,
                        user_id
                    ))

                conn.commit()

                return redirect(
                    url_for("profile")
                )

    except Exception:

        conn.rollback()
        raise

    finally:
        conn.close()

    return render_template(
        "edit_profile.html",
        user=user,
        error=error
    )


# ==================================================
# 출석체크
# ==================================================

@app.route(
    "/attendance",
    methods=["POST"]
)
def attendance():

    if not session.get(
        "user_logged_in"
    ):
        return redirect(
            url_for("login")
        )

    user_id = session.get(
        "user_id"
    )

    today = date.today()

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    point,
                    exp,
                    streak,
                    last_attendance
                FROM users
                WHERE id = %s
            """, (
                user_id,
            ))

            user = cur.fetchone()

            if user is None:

                return redirect(
                    url_for("login")
                )

            cur.execute("""
                SELECT id
                FROM attendance
                WHERE user_id = %s
                AND attendance_date = %s
            """, (
                user_id,
                today
            ))

            already_attended = cur.fetchone()

            if already_attended:

                return redirect(
                    url_for(
                        "profile",
                        attendance_message="already"
                    )
                )

            last_attendance = (
                user["last_attendance"]
            )

            if (
                last_attendance
                == today - timedelta(days=1)
            ):
                new_streak = (
                    user["streak"] + 1
                )

            else:
                new_streak = 1

            cur.execute("""
                INSERT INTO attendance (
                    user_id,
                    attendance_date
                )
                VALUES (
                    %s,
                    %s
                )
            """, (
                user_id,
                today
            ))

            cur.execute("""
                UPDATE users
                SET
                    point = point + 10,
                    exp = exp + 10,
                    streak = %s,
                    last_attendance = %s
                WHERE id = %s
            """, (
                new_streak,
                today,
                user_id
            ))

        conn.commit()

    except Exception:

        conn.rollback()
        raise

    finally:
        conn.close()

    return redirect(
        url_for(
            "profile",
            attendance_message="success"
        )
    )


# ==================================================
# 관리자 로그인 확인
# ==================================================

def admin_login_required(
    view_function
):

    @wraps(view_function)
    def wrapped_view(
        *args,
        **kwargs
    ):

        if not session.get(
            "admin_logged_in"
        ):
            return redirect(
                url_for(
                    "admin_login"
                )
            )

        return view_function(
            *args,
            **kwargs
        )

    return wrapped_view


# ==================================================
# 관리자 로그인
# ==================================================

@app.route(
    "/admin/login",
    methods=["GET", "POST"]
)
def admin_login():

    error = None

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        conn = get_db_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    SELECT *
                    FROM admins
                    WHERE username = %s
                """, (
                    username,
                ))

                admin_user = cur.fetchone()

        finally:
            conn.close()

        if (
            admin_user
            and check_password_hash(
                admin_user["password_hash"],
                password
            )
        ):

            session.clear()

            session["admin_logged_in"] = True
            session["admin_username"] = (
                admin_user["username"]
            )

            return redirect(
                url_for("admin")
            )

        error = (
            "아이디 또는 비밀번호가 올바르지 않습니다."
        )

    return render_template(
        "admin_login.html",
        error=error
    )


# ==================================================
# 관리자 로그아웃
# ==================================================

@app.route("/admin/logout")
def admin_logout():

    session.clear()

    return redirect(
        url_for("admin_login")
    )



# ==================================================
# 관리자 회원 관리 / 최신 엑셀 내보내기
# PostgreSQL을 원본으로 사용하므로 가입 및 정보 수정 내용이
# 별도의 파일 동기화 없이 다운로드 시 항상 반영됩니다.
# ==================================================

MEMBER_EXPORT_COLUMNS = [
    ("id", "회원번호"), ("name", "이름"), ("username", "아이디"),
    ("phone", "전화번호"), ("email", "이메일"),
    ("postal_code", "우편번호"), ("address", "배송지 주소"),
    ("address_detail", "상세 주소"), ("level", "레벨"),
    ("point", "포인트"), ("exp", "경험치(EXP)"),
    ("created_at", "가입일"),
]


def get_admin_members():
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""SELECT id, name, username, phone, email,
                           postal_code, address, address_detail,
                           level, point, exp, created_at
                           FROM users ORDER BY id DESC""")
            return cur.fetchall()
    finally:
        conn.close()


def get_member_admin_csrf():
    if "member_admin_csrf" not in session:
        session["member_admin_csrf"] = token_urlsafe(32)
    return session["member_admin_csrf"]


def verify_member_admin_csrf():
    from hmac import compare_digest
    expected = session.get("member_admin_csrf")
    supplied = request.form.get("csrf_token", "")
    if not expected or not compare_digest(expected, supplied):
        abort(403, "보안 토큰이 유효하지 않습니다.")


@app.route("/admin/members/<int:member_id>/edit", methods=["GET", "POST"])
@admin_login_required
def admin_member_edit(member_id):
    error = None
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""SELECT id, name, username, phone, email, postal_code,
                           address, address_detail, level, point, exp
                           FROM users WHERE id=%s""", (member_id,))
            member = cur.fetchone()
            if not member:
                abort(404)
            if request.method == "POST":
                verify_member_admin_csrf()
                name = request.form.get("name", "").strip()
                phone = request.form.get("phone", "").strip()
                email = request.form.get("email", "").strip()
                postal_code = request.form.get("postal_code", "").strip()
                address = request.form.get("address", "").strip()
                address_detail = request.form.get("address_detail", "").strip()
                try:
                    level = int(request.form.get("level", ""))
                    point = int(request.form.get("point", ""))
                except ValueError:
                    level, point = -1, -1
                if not name or len(name) > 100 or not email or len(email) > 254:
                    error = "이름과 이메일을 올바르게 입력해 주세요."
                elif len(phone) > 30 or (postal_code and not re.fullmatch(r"[0-9]{5}", postal_code)):
                    error = "전화번호 또는 우편번호를 확인해 주세요."
                elif len(address) > 300 or len(address_detail) > 300:
                    error = "주소는 각 300자 이내로 입력해 주세요."
                elif not 1 <= level <= 100 or not 0 <= point <= 1000000000:
                    error = "레벨은 1~100, 포인트는 0~10억 범위로 입력해 주세요."
                else:
                    cur.execute("SELECT 1 FROM users WHERE email=%s AND id<>%s LIMIT 1", (email, member_id))
                    if cur.fetchone():
                        error = "이미 다른 회원이 사용 중인 이메일입니다."
                    else:
                        cur.execute("""UPDATE users SET name=%s, phone=%s, email=%s,
                                       postal_code=%s, address=%s, address_detail=%s,
                                       level=%s, point=%s WHERE id=%s""",
                                    (name, phone, email, postal_code or None,
                                     address, address_detail, level, point, member_id))
                        conn.commit()
                        return redirect(url_for("admin_members"))
        return render_template("admin_member_edit.html", member=member,
                               error=error, member_csrf=get_member_admin_csrf())
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/admin/members/<int:member_id>/delete", methods=["POST"])
@admin_login_required
def admin_member_delete(member_id):
    verify_member_admin_csrf()
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM users WHERE id=%s FOR UPDATE", (member_id,))
            if not cur.fetchone():
                abort(404)
            # 주문 기록이 있는 회원은 삭제를 제한해 주문/환불/포인트 내역을 보존합니다.
            cur.execute("SELECT 1 FROM bank_orders WHERE user_id=%s LIMIT 1", (member_id,))
            if cur.fetchone():
                return ("주문 기록이 있는 회원은 관리자 화면에서 직접 삭제할 수 없습니다. "
                        "주문 및 환불 기록 보존을 위해 별도 탈퇴/비식별화 처리가 필요합니다.", 409)
            cur.execute("DELETE FROM attendance WHERE user_id=%s", (member_id,))
            cur.execute("DELETE FROM member_cart_items WHERE user_id=%s", (member_id,))
            cur.execute("DELETE FROM users WHERE id=%s", (member_id,))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return redirect(url_for("admin_members"))


@app.route("/admin/members")
@admin_login_required
def admin_members():
    members = get_admin_members()
    return render_template("admin_members.html", members=members, member_csrf=get_member_admin_csrf())


@app.route("/admin/members/export.xlsx")
@admin_login_required
def admin_members_export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    members = get_admin_members()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "회원 명단"
    sheet.append([label for _, label in MEMBER_EXPORT_COLUMNS])
    for member in members:
        values = []
        for key, _ in MEMBER_EXPORT_COLUMNS:
            value = member.get(key)
            # 사용자 입력을 엑셀 수식으로 해석하지 않도록 방지
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                value = "'" + value
            if key in ("phone", "postal_code") and value is not None:
                value = str(value)
            if isinstance(value, datetime) and value.tzinfo is not None:
                value = value.replace(tzinfo=None)
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(horizontal="center")
    widths = [12, 16, 20, 20, 32, 14, 45, 32, 12, 14, 16, 23]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for row in sheet.iter_rows(min_row=2):
        for index in (3, 5):
            row[index].number_format = "@"  # 전화번호 / 우편번호
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="MONO_members_" + date.today().strftime("%Y%m%d") + ".xlsx",
    )


# ==================================================
# 관리자 상품 등록
# ==================================================

@app.route(
    "/admin",
    methods=["GET", "POST"]
)
@admin_login_required
def admin():

    if request.method == "POST":

        # ------------------------------------------
        # 기본정보
        # ------------------------------------------

        name = request.form.get(
            "name",
            ""
        ).strip()

        category = request.form.get(
            "category",
            ""
        ).strip()

        image_url = request.form.get(
            "image_url",
            ""
        ).strip()

        description_html = request.form.get(
            "description_html",
            ""
        )

        smartstore_price = request.form.get(
            "smartstore_price",
            ""
        )

        smartstore_url = request.form.get(
            "smartstore_url",
            ""
        ).strip()

        coupang_price = request.form.get(
            "coupang_price",
            ""
        )

        coupang_url = request.form.get(
            "coupang_url",
            ""
        ).strip()

        bank_transfer_price = request.form.get(
            "bank_transfer_price", ""
        ).strip()


        # ------------------------------------------
        # 대표 이미지
        # URL 또는 Supabase 파일 업로드
        # ------------------------------------------

        image_file = request.files.get(
            "image_file"
        )

        if (
            image_file
            and image_file.filename
        ):

            try:

                uploaded_image_url = (
                    upload_product_image(
                        image_file,
                        "product-images"
                    )
                )

                # 파일을 선택한 경우 업로드된 이미지 우선
                image_url = uploaded_image_url

            except ValueError as e:

                return (
                    str(e),
                    400
                )

            except Exception as e:

                traceback.print_exc()

                return (
                    "상품 이미지 업로드 중 오류가 발생했습니다: "
                    + str(e),
                    500
                )


        # 로컬 파일 저장은 사용하지 않음
        image_file_name = None


        # ------------------------------------------
        # 상품 설명
        # HTML 또는 이미지
        # ------------------------------------------

        description_image_data = None
        description_image_mime = None

        description_file = request.files.get(
            "description_file"
        )

        if (
            description_file
            and description_file.filename
        ):

            if (
                description_file.mimetype
                and description_file.mimetype.startswith(
                    "image/"
                )
            ):

                description_image_data = (
                    description_file.read()
                )

                description_image_mime = (
                    description_file.mimetype
                )

                # 이미지가 선택되면 HTML은 비움
                description_html = ""

            else:

                try:

                    description_html = (
                        description_file
                        .read()
                        .decode("utf-8")
                    )

                except UnicodeDecodeError:

                    return (
                        "상품 설명 파일은 HTML 또는 이미지 파일만 사용할 수 있습니다.",
                        400
                    )


        # ------------------------------------------
        # 상품 저장
        # ------------------------------------------

        conn = get_db_connection()

        try:

            with conn.cursor() as cur:

                cur.execute("""
                    INSERT INTO products (
                        name,
                        category,
                        image_url,
                        image_file,
                        description_html,
                        description_image,
                        description_image_mime,
                        smartstore_price,
                        smartstore_url,
                        coupang_price,
                        coupang_url,
                        bank_transfer_price
                    )
                    VALUES (
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s
                    )
                    RETURNING id
                """, (
                    name,
                    category,
                    image_url,
                    image_file_name,
                    description_html,
                    description_image_data,
                    description_image_mime,

                    int(smartstore_price)
                    if smartstore_price
                    else None,

                    smartstore_url,

                    int(coupang_price)
                    if coupang_price
                    else None,

                    coupang_url,
                    int(bank_transfer_price) if bank_transfer_price else None
                ))

                product_id = (
                    cur.fetchone()["id"]
                )


            # --------------------------------------
            # 옵션
            # --------------------------------------

            option_names = (
                request.form.getlist(
                    "option_name[]"
                )
            )

            option_smartstore_prices = ["" for _ in option_names]

            option_smartstore_urls = ["" for _ in option_names]

            option_coupang_prices = ["" for _ in option_names]

            option_coupang_urls = ["" for _ in option_names]

            option_bank_transfer_prices = request.form.getlist(
                "option_bank_transfer_price[]"
            )



            for (
                option_name,
                option_smart_price,
                option_smart_url,
                option_coupang_price,
                option_coupang_url,
                option_bank_price
            ) in zip(
                option_names,
                option_smartstore_prices,
                option_smartstore_urls,
                option_coupang_prices,
                option_coupang_urls,
                option_bank_transfer_prices
            ):

                option_name = (
                    option_name.strip()
                )

                if not option_name:
                    continue

                with conn.cursor() as cur:

                    cur.execute("""
                        INSERT INTO product_options (
                            product_id,
                            option_name,
                            smartstore_price,
                            smartstore_url,
                            coupang_price,
                            coupang_url,
                            bank_transfer_price
                        )
                        VALUES (
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s
                        )
                    """, (
                        product_id,
                        option_name,

                        int(option_smart_price)
                        if option_smart_price
                        else None,

                        option_smart_url.strip(),

                        int(option_coupang_price)
                        if option_coupang_price
                        else None,

                        option_coupang_url.strip(),
                        int(option_bank_price) if option_bank_price else None
                    ))

            conn.commit()

        except Exception:

            conn.rollback()
            raise

        finally:

            conn.close()


        return redirect(
            url_for("admin")
        )


    # ----------------------------------------------
    # 관리자 페이지 상품 목록
    # ----------------------------------------------

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    name,
                    category,
                    image_url,
                    image_file,
                    description_html,

                    CASE
                        WHEN description_image IS NOT NULL
                        THEN TRUE
                        ELSE FALSE
                    END AS has_description_image,

                    smartstore_price,
                    smartstore_url,
                    coupang_price,
                    coupang_url,
                    bank_transfer_price,
                    is_sold_out,
                    created_at

                FROM products
                ORDER BY id DESC
            """)

            products = cur.fetchall()

    finally:

        conn.close()


    return render_template(
        "admin.html",
        products=products
    )


# ==================================================
# 상품 수정
# ==================================================

@app.route(
    "/admin/product/<int:product_id>/edit",
    methods=["GET", "POST"]
)
@admin_login_required
def edit_product(product_id):

    conn = get_db_connection()

    try:

        # ------------------------------------------
        # 현재 상품
        # ------------------------------------------

        with conn.cursor() as cur:

            cur.execute("""
                SELECT
                    id,
                    name,
                    category,
                    image_url,
                    image_file,
                    description_html,

                    CASE
                        WHEN description_image IS NOT NULL
                        THEN TRUE
                        ELSE FALSE
                    END AS has_description_image,

                    description_image_mime,
                    smartstore_price,
                    smartstore_url,
                    coupang_price,
                    coupang_url,
                    bank_transfer_price,
                    created_at

                FROM products
                WHERE id = %s
            """, (
                product_id,
            ))

            product = cur.fetchone()

            if product is None:

                return (
                    "상품을 찾을 수 없습니다.",
                    404
                )


            cur.execute("""
                SELECT *
                FROM product_options
                WHERE product_id = %s
                ORDER BY id ASC
            """, (
                product_id,
            ))

            options = cur.fetchall()


        # ------------------------------------------
        # 수정 저장
        # ------------------------------------------

        if request.method == "POST":

            name = request.form.get(
                "name",
                ""
            ).strip()

            category = request.form.get(
                "category",
                ""
            ).strip()

            image_url = request.form.get(
                "image_url",
                ""
            ).strip()

            description_html = request.form.get(
                "description_html",
                ""
            )

            smartstore_price = request.form.get(
                "smartstore_price",
                ""
            )

            smartstore_url = request.form.get(
                "smartstore_url",
                ""
            ).strip()

            coupang_price = request.form.get(
                "coupang_price",
                ""
            )

            coupang_url = request.form.get(
                "coupang_url",
                ""
            ).strip()

            bank_transfer_price = request.form.get(
                "bank_transfer_price", ""
            ).strip()


            # --------------------------------------
            # 대표 이미지 수정
            # --------------------------------------

            image_file = request.files.get(
                "image_file"
            )

            # URL 입력을 비워두었다면
            # 기존 이미지 유지
            if not image_url:
                image_url = product[
                    "image_url"
                ]


            # 새 이미지 파일을 선택했다면
            # Supabase에 업로드하고 새 URL 사용
            if (
                image_file
                and image_file.filename
            ):

                try:

                    uploaded_image_url = (
                        upload_product_image(
                            image_file,
                            "product-images"
                        )
                    )

                    image_url = (
                        uploaded_image_url
                    )

                except ValueError as e:

                    return (
                        str(e),
                        400
                    )

                except Exception as e:

                    traceback.print_exc()

                    return (
                        "상품 이미지 업로드 중 오류가 발생했습니다: "
                        + str(e),
                        500
                    )


            # 로컬 파일 저장은 사용하지 않음
            image_file_name = None


            # --------------------------------------
            # 상품 설명 수정
            # --------------------------------------

            description_file = (
                request.files.get(
                    "description_file"
                )
            )

            new_description_image = None
            new_description_mime = None

            use_new_description_image = False


            if (
                description_file
                and description_file.filename
            ):

                if (
                    description_file.mimetype
                    and description_file.mimetype.startswith(
                        "image/"
                    )
                ):

                    new_description_image = (
                        description_file.read()
                    )

                    new_description_mime = (
                        description_file.mimetype
                    )

                    use_new_description_image = True

                    # 새 설명 이미지가 있으면
                    # HTML 설명 제거
                    description_html = ""

                else:

                    try:

                        description_html = (
                            description_file
                            .read()
                            .decode("utf-8")
                        )

                    except UnicodeDecodeError:

                        return (
                            "상품 설명 파일은 HTML 또는 이미지 파일만 사용할 수 있습니다.",
                            400
                        )


            # --------------------------------------
            # 상품 정보 UPDATE
            # --------------------------------------

            with conn.cursor() as cur:

                if use_new_description_image:

                    cur.execute("""
                        UPDATE products
                        SET
                            name = %s,
                            category = %s,
                            image_url = %s,
                            image_file = %s,
                            description_html = %s,
                            description_image = %s,
                            description_image_mime = %s,
                            smartstore_price = %s,
                            smartstore_url = %s,
                            coupang_price = %s,
                            coupang_url = %s,
                            bank_transfer_price = %s
                        WHERE id = %s
                    """, (
                        name,
                        category,
                        image_url,
                        image_file_name,
                        description_html,
                        new_description_image,
                        new_description_mime,

                        int(smartstore_price)
                        if smartstore_price
                        else None,

                        smartstore_url,

                        int(coupang_price)
                        if coupang_price
                        else None,

                        coupang_url,
                        int(bank_transfer_price) if bank_transfer_price else None,
                        product_id
                    ))

                else:

                    cur.execute("""
                        UPDATE products
                        SET
                            name = %s,
                            category = %s,
                            image_url = %s,
                            image_file = %s,
                            description_html = %s,
                            smartstore_price = %s,
                            smartstore_url = %s,
                            coupang_price = %s,
                            coupang_url = %s,
                            bank_transfer_price = %s
                        WHERE id = %s
                    """, (
                        name,
                        category,
                        image_url,
                        image_file_name,
                        description_html,

                        int(smartstore_price)
                        if smartstore_price
                        else None,

                        smartstore_url,

                        int(coupang_price)
                        if coupang_price
                        else None,

                        coupang_url,
                        int(bank_transfer_price) if bank_transfer_price else None,
                        product_id
                    ))


                # 기존 옵션 삭제
                cur.execute("""
                    DELETE FROM product_options
                    WHERE product_id = %s
                """, (
                    product_id,
                ))


            # --------------------------------------
            # 옵션 다시 저장
            # --------------------------------------

            option_names = (
                request.form.getlist(
                    "option_name[]"
                )
            )

            option_smartstore_prices = ["" for _ in option_names]

            option_smartstore_urls = ["" for _ in option_names]

            option_coupang_prices = ["" for _ in option_names]

            option_coupang_urls = ["" for _ in option_names]

            option_bank_transfer_prices = request.form.getlist(
                "option_bank_transfer_price[]"
            )



            for (
                option_name,
                smart_price,
                smart_url,
                cp_price,
                cp_url,
                bank_price
            ) in zip(
                option_names,
                option_smartstore_prices,
                option_smartstore_urls,
                option_coupang_prices,
                option_coupang_urls,
                option_bank_transfer_prices
            ):

                option_name = (
                    option_name.strip()
                )

                if not option_name:
                    continue


                with conn.cursor() as cur:

                    cur.execute("""
                        INSERT INTO product_options (
                            product_id,
                            option_name,
                            smartstore_price,
                            smartstore_url,
                            coupang_price,
                            coupang_url,
                            bank_transfer_price
                        )
                        VALUES (
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s
                        )
                    """, (
                        product_id,
                        option_name,

                        int(smart_price)
                        if smart_price
                        else None,

                        smart_url.strip(),

                        int(cp_price)
                        if cp_price
                        else None,

                        cp_url.strip(),
                        int(bank_price) if bank_price else None
                    ))


            conn.commit()

            return redirect(
                url_for("admin")
            )


        return render_template(
            "edit_product.html",
            product=product,
            options=options
        )


    except Exception:

        conn.rollback()
        raise

    finally:

        conn.close()


# ==================================================
# 상품 품절 / 판매 재개
# ==================================================

@app.route(
    "/admin/product/<int:product_id>/sold-out",
    methods=["POST"]
)
@admin_login_required
def toggle_sold_out(product_id):

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            # 현재 품절 상태 확인
            cur.execute("""
                SELECT is_sold_out
                FROM products
                WHERE id = %s
            """, (
                product_id,
            ))

            product = cur.fetchone()

            if product is None:

                return (
                    "상품을 찾을 수 없습니다.",
                    404
                )

            # 현재 상태 반대로 변경
            new_status = not bool(
                product["is_sold_out"]
            )

            cur.execute("""
                UPDATE products
                SET is_sold_out = %s
                WHERE id = %s
            """, (
                new_status,
                product_id
            ))

        conn.commit()

    except Exception:

        conn.rollback()
        raise

    finally:

        conn.close()

    return redirect(
        url_for("admin")
    )


# ==================================================
# 상품 삭제
# ==================================================

@app.route(
    "/admin/product/<int:product_id>/delete",
    methods=["POST"]
)
@admin_login_required
def delete_product(product_id):

    conn = get_db_connection()

    try:

        with conn.cursor() as cur:

            cur.execute("""
                DELETE FROM products
                WHERE id = %s
            """, (
                product_id,
            ))

        conn.commit()

    except Exception:

        conn.rollback()
        raise

    finally:

        conn.close()


    return redirect(
        url_for("admin")
    )


# ==================================================
# 장바구니 / 바로 결제 (서버에서 가격 재확인)
# ==================================================

def cart_lines():
    """로그인한 회원 본인의 장바구니만 조회한다. 가격은 DB에서 재검증한다."""
    user_id = session.get("user_id")
    if not session.get("user_logged_in") or not user_id:
        return []
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""SELECT c.id AS cart_id, c.product_id, c.option_id, c.quantity,
                p.name AS product_name, p.image_url,
                CASE WHEN c.option_id IS NULL THEN '기본 상품' ELSE o.option_name END AS option_name,
                CASE WHEN c.option_id IS NULL THEN p.bank_transfer_price ELSE o.bank_transfer_price END AS price
                FROM member_cart_items c
                JOIN products p ON p.id = c.product_id
                LEFT JOIN product_options o ON o.id = c.option_id AND o.product_id = c.product_id
                WHERE c.user_id = %s AND p.is_sold_out = FALSE
                AND (c.option_id IS NULL OR o.id IS NOT NULL)
                ORDER BY c.id ASC LIMIT 30""", (user_id,))
            rows = cur.fetchall()
        return [dict(product_id=r['product_id'], option_id=r['option_id'],
                     product_name=r['product_name'], option_name=r['option_name'],
                     image_url=r['image_url'], price=r['price'], quantity=r['quantity'],
                     subtotal=r['price']*r['quantity'])
                for r in rows if r['price'] is not None and r['price'] > 0]
    finally:
        conn.close()


@app.route("/cart")
def view_cart():
    lines = cart_lines()
    return render_template("cart.html", lines=lines,
                           total=sum(x["subtotal"] for x in lines), csrf=get_order_csrf())


@app.route("/cart/add", methods=["POST"])
def add_to_cart():
    if request.form.get("csrf_token") != session.get("order_csrf") or not session.get("order_csrf"):
        abort(403)
    try:
        pid = int(request.form["product_id"])
        oid = int(request.form["option_id"]) if request.form.get("option_id") else None
        qty = int(request.form.get("quantity", 1))
        if not (1 <= qty <= 20):
            raise ValueError()
    except (ValueError, KeyError):
        abort(400)
    user_id = session.get("user_id")
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT bank_transfer_price, is_sold_out FROM products WHERE id=%s", (pid,))
            product = cur.fetchone()
            if not product or product['is_sold_out']:
                abort(400, "판매 중인 상품이 아닙니다.")
            price = product['bank_transfer_price']
            if oid is not None:
                cur.execute("SELECT bank_transfer_price FROM product_options WHERE id=%s AND product_id=%s", (oid, pid))
                option = cur.fetchone()
                if not option:
                    abort(400, "옵션을 찾을 수 없습니다.")
                price = option['bank_transfer_price']
            if price is None or price <= 0:
                abort(400, "판매가격이 없는 상품입니다.")
            cur.execute("SELECT COUNT(*) AS n FROM member_cart_items WHERE user_id=%s", (user_id,))
            count = cur.fetchone()['n']
            cur.execute("SELECT id FROM member_cart_items WHERE user_id=%s AND product_id=%s AND option_id IS NOT DISTINCT FROM %s", (user_id,pid,oid))
            existing = cur.fetchone()
            if not existing and count >= 30:
                abort(400, "장바구니는 최대 30개 상품까지 담을 수 있습니다.")
            cur.execute("""INSERT INTO member_cart_items (user_id, product_id, option_id, quantity)
                VALUES (%s,%s,%s,%s)
                ON CONFLICT (user_id, product_id, (COALESCE(option_id, 0)))
                DO UPDATE SET quantity = LEAST(20, member_cart_items.quantity + EXCLUDED.quantity)""",
                (user_id, pid, oid, qty))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return redirect(url_for("checkout") if request.form.get("mode") == "buy" else url_for("view_cart"))


@app.route("/cart/update", methods=["POST"])
def update_cart():
    if request.form.get("csrf_token") != session.get("order_csrf") or not session.get("order_csrf"):
        abort(403)
    try:
        index = int(request.form.get("index"))
        quantity = int(request.form.get("quantity"))
    except (TypeError, ValueError):
        abort(400)
    if not 0 <= quantity <= 20:
        abort(400)
    user_id = session.get("user_id")
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            # cart.html은 화면에 표시된 행의 0부터 시작하는 인덱스를 제출한다.
            # 조회 시 숨겨지는 품절/가격 미등록 항목은 제외해 같은 순서를 유지한다.
            cur.execute("""SELECT c.id FROM member_cart_items c
                JOIN products p ON p.id=c.product_id
                LEFT JOIN product_options o ON o.id=c.option_id AND o.product_id=c.product_id
                WHERE c.user_id=%s AND p.is_sold_out=FALSE
                AND (c.option_id IS NULL OR o.id IS NOT NULL)
                AND (CASE WHEN c.option_id IS NULL THEN p.bank_transfer_price ELSE o.bank_transfer_price END)>0
                ORDER BY c.id ASC LIMIT 30""", (user_id,))
            rows = cur.fetchall()
            if not 0 <= index < len(rows):
                abort(400)
            cart_id = rows[index]['id']
            if quantity == 0:
                cur.execute("DELETE FROM member_cart_items WHERE id=%s AND user_id=%s", (cart_id, user_id))
            else:
                cur.execute("UPDATE member_cart_items SET quantity=%s WHERE id=%s AND user_id=%s", (quantity,cart_id,user_id))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return redirect(url_for("view_cart"))


@app.route("/checkout", methods=["GET", "POST"])
def checkout():
    user_id = session.get("user_id")
    if not session.get("user_logged_in") or not user_id:
        return redirect(url_for("login"))
    lines = cart_lines()
    if not lines:
        return redirect(url_for("view_cart"))
    total = sum(x["subtotal"] for x in lines)
    error = None
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""SELECT id, name, phone, postal_code, address, address_detail, point
                           FROM users WHERE id=%s""", (user_id,))
            user = cur.fetchone()
    finally:
        conn.close()
    if not user:
        return redirect(url_for("login"))
    requested_points = 0
    if request.method == "POST":
        if request.form.get("csrf_token") != session.get("order_csrf") or not session.get("order_csrf"):
            abort(403)
        depositor = request.form.get("depositor_name", "").strip()
        recipient_name = request.form.get("recipient_name", "").strip()
        recipient_phone = request.form.get("recipient_phone", "").strip()
        postal_code = request.form.get("postal_code", "").strip()
        address = request.form.get("address", "").strip()
        address_detail = request.form.get("address_detail", "").strip()
        raw_points = request.form.get("use_points", "0").strip() or "0"
        if not re.fullmatch(r"[0-9]{1,12}", raw_points):
            error = "사용 포인트는 0 이상의 정수로 입력해 주세요."
        else:
            requested_points = int(raw_points)
        if not error and requested_points > min(total, max(0, user["point"] or 0)):
            error = "사용 가능한 포인트를 초과했습니다."
        if not error and not BANK_ACCOUNT_NUMBER and requested_points < total:
            error = "계좌번호가 등록되지 않아 주문할 수 없습니다."
        if not error and not 1 <= len(depositor) <= 50 and requested_points < total:
            error = "입금자명을 1~50자로 입력해 주세요."
        elif not error and not 1 <= len(recipient_name) <= 100:
            error = "받는 분 이름을 입력해 주세요."
        elif not error and not 1 <= len(recipient_phone) <= 30:
            error = "받는 분 연락처를 입력해 주세요."
        elif not error and not re.fullmatch(r"[0-9]{5}", postal_code):
            error = "우편번호 5자리를 입력해 주세요."
        elif not error and (not address or len(address) > 300 or len(address_detail) > 300):
            error = "배송지 주소를 정확히 입력해 주세요."
        if not error:
            conn = get_db_connection()
            try:
                with conn.cursor() as cur:
                    # 회원 행을 잠가 동시 주문/포인트 중복 차감을 방지한다.
                    cur.execute("SELECT point FROM users WHERE id=%s FOR UPDATE", (user_id,))
                    locked_user = cur.fetchone()
                    if not locked_user:
                        abort(403)
                    # 동일 회원의 장바구니 행을 잠근 뒤 서버에서 가격을 재조회한다.
                    cur.execute("""SELECT c.id, c.product_id, c.option_id, c.quantity,
                        p.name AS product_name,
                        CASE WHEN c.option_id IS NULL THEN '기본 상품' ELSE o.option_name END AS option_name,
                        CASE WHEN c.option_id IS NULL THEN p.bank_transfer_price ELSE o.bank_transfer_price END AS price
                        FROM member_cart_items c JOIN products p ON p.id=c.product_id
                        LEFT JOIN product_options o ON o.id=c.option_id AND o.product_id=c.product_id
                        WHERE c.user_id=%s AND p.is_sold_out=FALSE
                        AND (c.option_id IS NULL OR o.id IS NOT NULL)
                        ORDER BY c.id ASC LIMIT 30 FOR UPDATE OF c""", (user_id,))
                    live_rows = cur.fetchall()
                    if not live_rows or any(r['price'] is None or r['price'] <= 0 for r in live_rows):
                        error = "장바구니 상품 정보가 변경되었습니다. 다시 확인해 주세요."
                    else:
                        live_total = sum(r['price'] * r['quantity'] for r in live_rows)
                        if live_total != total or len(live_rows) != len(lines):
                            error = "주문 금액이 변경되었습니다. 장바구니를 다시 확인해 주세요."
                        elif requested_points > min(live_total, max(0, locked_user['point'] or 0)):
                            error = "포인트 잔액이 변경되었습니다. 다시 확인해 주세요."
                        elif not BANK_ACCOUNT_NUMBER and requested_points < live_total:
                            error = "계좌번호가 등록되지 않아 주문할 수 없습니다."
                        else:
                            # 주문별 포인트는 합계가 정확히 일치하도록 순서대로 배분한다.
                            remaining_points = requested_points
                            group_id = uuid4()
                            ids = []
                            for i, row in enumerate(live_rows):
                                gross = row['price'] * row['quantity']
                                applied = min(gross, remaining_points)
                                remaining_points -= applied
                                cash = gross - applied
                                order_status = 'paid' if cash == 0 else 'reported'
                                cur.execute("""INSERT INTO bank_orders
                                    (product_id, product_name, option_name, depositor_name, quantity,
                                     amount, gross_amount, point_used, order_group_id, status, paid_at,
                                     recipient_name, recipient_phone, postal_code, address, address_detail, user_id)
                                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,
                                            CASE WHEN %s='paid' THEN NOW() ELSE NULL END,
                                            %s,%s,%s,%s,%s,%s) RETURNING id""",
                                    (row['product_id'], row['product_name'],
                                     row['option_name'] + (' × ' + str(row['quantity']) if row['quantity'] > 1 else ''),
                                     depositor or '포인트 전액 결제', row['quantity'], cash, gross, applied, group_id,
                                     order_status, order_status, recipient_name, recipient_phone,
                                     postal_code, address, address_detail, user_id))
                                ids.append(cur.fetchone()['id'])
                            if requested_points:
                                cur.execute("UPDATE users SET point=point-%s WHERE id=%s", (requested_points, user_id))
                            cur.execute("DELETE FROM member_cart_items WHERE user_id=%s", (user_id,))
                if error:
                    conn.rollback()
                else:
                    conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()
            if not error:
                session['order_csrf'] = token_urlsafe(32)
                return render_template("checkout_success.html", ids=ids,
                                       total=total-requested_points)
    return render_template("checkout.html", lines=lines, total=total,
                           csrf=get_order_csrf(), account=BANK_ACCOUNT_NUMBER,
                           error=error, user=user, requested_points=requested_points)


# ==================================================
# 계좌이체 주문 접수 / 관리자 주문 및 매출 관리
# ==================================================

@app.route("/api/bank-orders", methods=["POST"])
def create_bank_order():
    if not BANK_ACCOUNT_NUMBER:
        return jsonify(error="계좌번호가 등록되지 않아 주문을 받을 수 없습니다."), 503
    data = request.get_json(silent=True) or {}
    if data.get("csrf_token") != session.get("order_csrf") or not session.get("order_csrf"):
        return jsonify(error="보안 토큰이 유효하지 않습니다. 페이지를 새로고침해 주세요."), 403
    depositor = str(data.get("depositor_name") or "").strip()
    if not (1 <= len(depositor) <= 50):
        return jsonify(error="입금자명을 1~50자로 입력해 주세요."), 400
    try:
        product_id = int(data.get("product_id"))
    except (TypeError, ValueError):
        return jsonify(error="상품 정보가 올바르지 않습니다."), 400
    option_id = data.get("option_id")
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id,name,bank_transfer_price,is_sold_out FROM products WHERE id=%s", (product_id,))
            product = cur.fetchone()
            if not product or product["is_sold_out"]:
                return jsonify(error="판매 중인 상품이 아닙니다."), 400
            option_name = "기본 상품"
            amount = product["bank_transfer_price"]
            if option_id is not None:
                try:
                    option_id = int(option_id)
                except (TypeError, ValueError):
                    return jsonify(error="옵션 정보가 올바르지 않습니다."), 400
                cur.execute("SELECT option_name,bank_transfer_price FROM product_options WHERE id=%s AND product_id=%s", (option_id, product_id))
                option = cur.fetchone()
                if not option:
                    return jsonify(error="옵션을 찾을 수 없습니다."), 400
                option_name, amount = option["option_name"], option["bank_transfer_price"]
            if amount is None or amount <= 0:
                return jsonify(error="계좌이체 판매가격이 등록되지 않았습니다."), 400
            cur.execute("""INSERT INTO bank_orders (product_id,product_name,option_name,depositor_name,amount)
                           VALUES (%s,%s,%s,%s,%s) RETURNING id""",
                        (product_id, product["name"], option_name, depositor, amount))
            order_id = cur.fetchone()["id"]
        conn.commit()
        return jsonify(ok=True, order_id=order_id, message="입금 신고가 접수되었습니다. 관리자가 실제 입금을 확인한 뒤 결제 완료로 처리합니다."), 201
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@app.route("/admin/orders")
@admin_login_required
def admin_orders():
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT * FROM bank_orders ORDER BY created_at DESC LIMIT 500")
            orders = cur.fetchall()
    finally:
        conn.close()
    return render_template("admin_orders.html", orders=orders, csrf_token=get_order_csrf())


@app.route("/admin/orders/<int:order_id>/status", methods=["POST"])
@admin_login_required
def admin_order_status(order_id):
    if request.form.get("csrf_token") != session.get("order_csrf") or not session.get("order_csrf"):
        abort(403)
    status = request.form.get("status")
    if status not in ("paid", "cancelled", "refunded"):
        abort(400)
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT status, user_id, point_used FROM bank_orders WHERE id=%s FOR UPDATE", (order_id,))
            order = cur.fetchone()
            if not order:
                abort(404)
            old = order["status"]
            allowed = {"reported": {"paid", "cancelled"}, "paid": {"refunded"}}
            if status not in allowed.get(old, set()):
                abort(400, "허용되지 않는 주문 상태 변경입니다.")
            if status in ('cancelled', 'refunded') and order['user_id'] and order['point_used']:
                cur.execute("UPDATE users SET point=point+%s WHERE id=%s",
                            (order['point_used'], order['user_id']))
            cur.execute("""UPDATE bank_orders SET status=%s,
                paid_at=CASE WHEN %s='paid' THEN NOW() ELSE paid_at END,
                refunded_at=CASE WHEN %s='refunded' THEN NOW() ELSE refunded_at END
                WHERE id=%s""", (status, status, status, order_id))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return redirect(url_for("admin_orders"))


@app.route("/admin/sales")
@admin_login_required
def admin_sales():
    # 결제확인 시점 기준, 환불 주문은 전체 기간 집계에서 제외합니다.
    today = date.today()
    start = request.args.get("start", "")
    end = request.args.get("end", "")
    try:
        start_date = date.fromisoformat(start) if start else today.replace(day=1)
        end_date = date.fromisoformat(end) if end else today
        if start_date > end_date or (end_date - start_date).days > 3660:
            raise ValueError()
    except ValueError:
        abort(400, "올바른 조회 기간을 입력해 주세요.")
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("""SELECT
                COALESCE(SUM(amount) FILTER (WHERE (paid_at AT TIME ZONE 'Asia/Seoul')::date=%s),0) AS today_sales,
                COALESCE(SUM(amount) FILTER (WHERE date_trunc('month',paid_at AT TIME ZONE 'Asia/Seoul')=date_trunc('month',NOW() AT TIME ZONE 'Asia/Seoul')),0) AS month_sales,
                COALESCE(SUM(amount),0) AS total_sales, COUNT(*) AS paid_count
                FROM bank_orders WHERE status='paid'""", (today,))
            summary = cur.fetchone()
            cur.execute("""SELECT (paid_at AT TIME ZONE 'Asia/Seoul')::date AS day,
                SUM(amount) AS amount,COUNT(*) AS orders FROM bank_orders
                WHERE status='paid' AND (paid_at AT TIME ZONE 'Asia/Seoul')::date BETWEEN %s AND %s
                GROUP BY 1 ORDER BY 1""", (start_date, end_date))
            daily = cur.fetchall()
            cur.execute("""SELECT product_name, SUM(amount) AS amount,COUNT(*) AS orders
                FROM bank_orders WHERE status='paid' AND (paid_at AT TIME ZONE 'Asia/Seoul')::date BETWEEN %s AND %s
                GROUP BY product_name ORDER BY amount DESC LIMIT 20""", (start_date, end_date))
            ranking = cur.fetchall()
            cur.execute("""SELECT * FROM bank_orders WHERE status='paid'
                AND (paid_at AT TIME ZONE 'Asia/Seoul')::date BETWEEN %s AND %s
                ORDER BY paid_at DESC LIMIT 100""", (start_date, end_date))
            recent = cur.fetchall()
    finally:
        conn.close()
    return render_template("admin_sales.html", summary=summary, daily=daily,
                           ranking=ranking, recent=recent, start_date=start_date, end_date=end_date)


# ==================================================
# 로컬 실행
# ==================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )
